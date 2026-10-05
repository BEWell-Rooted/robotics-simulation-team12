"""Presentation deck for assignment 1 -> docs/assignment1_slides_v1.pptx (text: docs/slides_draft_v0.md).

Same visual system as the interim report page (docs/interim_report_1001.md web version, "Labnote"): warm paper
ground, oxblood section marks, hairline rules, serif titles, mono numbers.
  fonts : titles "Noto Serif KR" Bold · body "Pretendard" · numbers "JetBrains Mono" (all free; install before
          presenting, or embed them when saving from PowerPoint)
  colors: ink #241e19 / #55493f / #756558, paper #faf6f0, brand (oxblood) #7a2e2e, accent #a6741b, data #2e5163

    <venv with python-pptx>/bin/python scripts/make_slides.py
"""

import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PROJECT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(PROJECT, "docs", "figures")
OUT = os.path.join(PROJECT, "docs", "assignment1_slides_v1.pptx")
N_SLIDES = 12

SERIF, SANS, MONO = "Noto Serif KR", "Pretendard", "JetBrains Mono"
C = {k: RGBColor.from_string(v) for k, v in dict(
    ink900="241E19", ink700="55493F", ink600="756558", ink500="9A8877", ink300="D8CDBF", ink200="E9E1D6",
    ink100="F3EDE4", ink050="FAF6F0", white="FFFFFF", brand="7A2E2E", brand_soft="F7E9E8", accent="A6741B",
    accent_soft="F6ECD6", data="2E5163", data_soft="DFE9EE", done="2F6B4F", blocked="B4432F").items()}  # fmt: skip
W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.75)  # side margin
CW = W - 2 * MX  # content width


# ------------------------------------------------------------------------------------------------ text primitives


def font(run, size, face=SANS, bold=False, color="ink900", spacing=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = C[color]
    run.font.name = face
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", SANS if face == MONO else face)  # Korean glyphs of mono runs come from the body font
    if spacing is not None:
        rpr.set("spc", str(spacing))


TOKEN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def runs(par, text, size, face=SANS, color="ink900", bold=False):
    """Inline markup: **bold**, `mono` (data color)."""
    for part in TOKEN.split(text):
        if not part:
            continue
        r = par.add_run()
        if part.startswith("**"):
            r.text = part[2:-2]
            font(r, size, face, True, color)
        elif part.startswith("`"):
            r.text = part[1:-1]
            font(r, size * 0.92, MONO, bold, "data")
        else:
            r.text = part
            font(r, size, face, bold, color)


def bullet(par, level):
    """Real hanging-indent bullets (• ink500, – for level 2)."""
    ppr = par._p.get_or_add_pPr()
    indent = Inches(0.24)
    ppr.set("marL", str(int(indent * level)))
    ppr.set("indent", str(-int(indent)))
    for tag, attrs in (("a:buClr", None), ("a:buFont", {"typeface": "Arial"}), ("a:buChar", {"char": "•" if level == 1 else "–"})):
        el = ppr.makeelement(qn(tag), attrs or {})
        if tag == "a:buClr":
            clr = el.makeelement(qn("a:srgbClr"), {"val": "9A8877" if level == 1 else "D8CDBF"})
            el.append(clr)
        ppr.append(el)


def text(slide, x, y, w, h, lines, size=17, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, gap=7, line=1.18):
    """lines: str (bullet) or dict(t=, lv=0|1|2, size=, face=, color=, bold=, gap=)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        ln = {"t": ln, "lv": 1} if isinstance(ln, str) else ln
        lv = ln.get("lv", 0)
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.line_spacing = line
        par.space_after = Pt(ln.get("gap", gap))
        if lv:
            bullet(par, lv)
        sz = ln.get("size", size - (2 if lv == 2 else 0))
        runs(par, ln["t"], sz, ln.get("face", SANS), ln.get("color", "ink700" if lv == 2 else "ink900"), ln.get("bold", False))
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    eff = s._element.find(".//" + qn("a:effectRef"))
    if eff is not None:
        eff.set("idx", "0")  # no theme effect (LibreOffice draws the theme shadow otherwise)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = C[fill]
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = C[line]
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    return s


def hline(slide, x, y, w, color="ink200", weight=0.75):
    return rect(slide, x, y, w, Pt(weight), fill=color)


# ------------------------------------------------------------------------------------------------ page furniture


def page(prs, n, kicker, title_text, source=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C["ink050"]
    hline(s, MX, Inches(0.42), CW, "brand", 2.25)  # h2 border-top
    text(s, MX, Inches(0.55), CW, Inches(0.3), [dict(t=f"{n:02d} · {kicker}", size=11, color="brand", gap=0)])
    t = text(s, MX, Inches(0.8), CW, Inches(0.7), [dict(t=title_text, size=27, face=SERIF, bold=True, gap=0)])
    t.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("spc", "-30")
    hline(s, MX, H - Inches(0.55), CW)
    if source:
        text(s, MX, H - Inches(0.45), CW - Inches(1.2), Inches(0.3), [dict(t=source, size=10, color="ink600", gap=0)])
    text(s, W - MX - Inches(1.0), H - Inches(0.45), Inches(1.0), Inches(0.3),
         [dict(t=f"{n:02d} / {N_SLIDES}", size=10, face=MONO, color="ink500", gap=0)], align=PP_ALIGN.RIGHT)  # fmt: skip
    return s


def kicker_line(slide, x, y, w, label, color="brand"):
    text(slide, x, y, w, Inches(0.3), [dict(t=label, size=11, color=color, gap=0)])


def callout(slide, x, y, w, h, label, lines, tone="brand", size=15):
    rect(slide, x, y, w, h, fill=f"{tone}_soft")
    rect(slide, x, y, Pt(2.5), h, fill=tone)
    kicker_line(slide, x + Inches(0.22), y + Inches(0.14), w - Inches(0.4), label, tone)
    text(slide, x + Inches(0.22), y + Inches(0.44), w - Inches(0.4), h - Inches(0.5), lines, size=size, gap=5)


def pill(slide, x, y, label, tone, size=13):
    """Coloured dot + label (done = 지지, blocked = 기각, accent = 부분)."""
    rect(slide, x, y + Inches(0.07), Inches(0.11), Inches(0.11), fill=tone, shape=MSO_SHAPE.OVAL)
    text(slide, x + Inches(0.18), y, Inches(3.5), Inches(0.3), [dict(t=label, size=size, color=tone, bold=True, gap=0)])


def stats(slide, x, y, w, items, size=28):
    """[(value, label, brand?)] -> big mono numbers with a 2px ink rule on top."""
    n, gap = len(items), Inches(0.3)
    cw = (w - gap * (n - 1)) / n
    for i, (v, label, brand) in enumerate(items):
        cx = x + i * (cw + gap)
        hline(slide, cx, y, cw, "ink900", 1.5)
        text(slide, cx, y + Inches(0.12), cw, Inches(0.55), [dict(t=v, size=size, face=MONO, color="brand" if brand else "ink900", gap=0)])
        text(slide, cx, y + Inches(0.7), cw, Inches(0.6), [dict(t=label, size=12, color="ink600", gap=0)], line=1.1)


def figure(slide, name, x, y, w, caption=None, num=None):
    from PIL import Image  # noqa: PLC0415

    iw, ih = Image.open(os.path.join(FIG, name)).size
    pad = Inches(0.08)
    ph = (w - 2 * pad) * ih / iw
    rect(slide, x, y, w, ph + 2 * pad, fill="white", line="ink200")
    slide.shapes.add_picture(os.path.join(FIG, name), x + pad, y + pad, width=w - 2 * pad)
    if caption:
        text(slide, x, y + ph + 2 * pad + Inches(0.08), w, Inches(0.5),
             [dict(t=(f"**그림 {num}.** " if num else "") + caption, size=11, color="ink600", gap=0)], line=1.1)  # fmt: skip
    return ph + 2 * pad


def cell_border(cell, side, color, width_pt):
    tcpr = cell._tc.get_or_add_tcPr()
    tag = {"B": "a:lnB", "T": "a:lnT", "L": "a:lnL", "R": "a:lnR"}[side]
    for old in tcpr.findall(qn(tag)):
        tcpr.remove(old)
    ln = tcpr.makeelement(qn(tag), {"w": str(int(Pt(width_pt))), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    if color is None:
        ln.append(ln.makeelement(qn("a:noFill"), {}))
    else:
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(C[color])}))
        ln.append(fill)
    tcpr.append(ln)


def table(slide, x, y, w, rows, col_w, size=14, num_cols=(), hl=(), best=(), row_h=0.42):
    """Report-style table: header ink700 with 2px ink900 rule, 1px ink200 row rules, mono right-aligned numbers.
    hl = highlighted rows (brand_soft), best = (row, col) cells in brand bold."""
    gf = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
    tbl_pr = gf._element.graphic.graphicData.tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tbl_pr.set(attr, "0")
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is None:
        style = tbl_pr.makeelement(qn("a:tableStyleId"), {})
        tbl_pr.append(style)
    style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    tbl = gf.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Emu(int(w * cw / sum(col_w)))
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if i in hl:
                cell.fill.solid()
                cell.fill.fore_color.rgb = C["brand_soft"]
            else:
                cell.fill.background()
            for side in ("L", "R", "T"):
                cell_border(cell, side, None, 0)
            cell_border(cell, "B", "ink900" if i == 0 else "ink200", 1.5 if i == 0 else 0.75)
            par = cell.text_frame.paragraphs[0]
            num = j in num_cols
            par.alignment = PP_ALIGN.RIGHT if num else PP_ALIGN.LEFT
            if i == 0:
                runs(par, val, size - 1, color="ink700")
            else:
                is_best = (i, j) in best
                runs(par, val, size * (0.92 if num else 1), face=MONO if num else SANS,
                     color="brand" if is_best else "ink900", bold=is_best or (i in hl and not num))  # fmt: skip
    return gf


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------------------------------------ slides


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1. cover ------------------------------------------------------------------------------------------------------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C["ink050"]
    sheet_x, sheet_y = Inches(0.9), Inches(0.7)
    rect(s, sheet_x, sheet_y, W - 2 * sheet_x, H - 2 * sheet_y, fill="white")
    ix = sheet_x + Inches(0.7)
    iw = W - 2 * ix
    kicker_line(s, ix, Inches(1.35), iw, "ROBOTICS SIMULATION · 과제 1 · 팀 12")
    text(s, ix, Inches(1.7), iw, Inches(1.0), [dict(t="처음 보는 지형에서도 걷는 Ant", size=40, face=SERIF, bold=True, gap=0)])
    text(s, ix, Inches(2.75), iw, Inches(0.9), [dict(
        t="평면에서만 배운 Ant는 왜 처음 보는 지형에서 무너지는가 — 관찰에서 가설 5개를 세우고, 통제 실험과 잠금 테스트로 검증했다.",
        size=17, color="ink700", gap=0)], line=1.3)  # fmt: skip
    meta = "**팀 12** 박형근 · 경한 · [구성원]      **발표** 2026.10.08      **환경** Isaac Lab 2.3.0 · RSL-RL PPO"
    text(s, ix, Inches(3.75), iw, Inches(0.35), [dict(t=meta, size=13, color="ink600", gap=0)])
    hline(s, ix, Inches(4.2), iw)
    stats(s, ix, Inches(4.6), iw, [
        ("91.2", "처음 보는 지형 12종 평균 reward\nbaseline 16.3 → 최종 V8", True),
        ("113.65", "잠금 테스트 UnseenMix (1회)\nbaseline 9.51", False),
        ("174.73", "제출 태스크 공식 평면\nbaseline 134.90", False),
        ("5", "가설 — 지지 3 · 부분 1 · 기각 1", False),
    ], size=26)  # fmt: skip
    notes(s, "팀 소개. \"평면에서만 배운 Ant를 처음 보는 지형에서도 걷게 만든 과정을 가설과 검증 순서로 말씀드리겠습니다.\"\n"
             "[구성원] 이름 채우기.")  # fmt: skip

    # 2. problem ----------------------------------------------------------------------------------------------------
    s = page(prs, 2, "연구 문제", "왜 처음 보는 지형에서 무너지는가",
             "출처: docs/report_material.md 2절 (진단 실험, log_0930)")  # fmt: skip
    callout(s, MX, Inches(1.75), CW, Inches(1.05), "연구 문제",
            [dict(t="평면에서 학습한 Ant는 왜 처음 보는 지형에서 무너지는가?", size=21, face=SERIF, bold=True, color="brand")])  # fmt: skip
    col = (CW - Inches(0.5)) / 2
    for i, (lab, v, d) in enumerate([
        ("관찰 1 · 성능 붕괴", "134.90 → 16.3", "baseline 공식 평면 → 처음 보는 지형 12종 평균"),
        ("관찰 2 · 실패 = 뒤집힘", "59 m → 59 m", "종료 높이 임계 0.31 → 0.20에도 전진 거리 그대로 — 이미 뒤집혀 누운 상태"),
    ]):  # fmt: skip
        cx = MX + i * (col + Inches(0.5))
        hline(s, cx, Inches(3.15), col, "ink900", 1.5)
        kicker_line(s, cx, Inches(3.27), col, lab, "ink700")
        text(s, cx, Inches(3.6), col, Inches(0.6), [dict(t=v, size=30, face=MONO, color="brand", gap=0)])
        text(s, cx, Inches(4.3), col, Inches(0.7), [dict(t=d, size=14, color="ink600", gap=0)], line=1.2)
    text(s, MX, Inches(5.25), CW, Inches(1.5), [
        "과제: Isaac-Ant-v0 정책이 **지형 형태·마찰만 바뀐 처음 보는 환경**에서도 걷게 만들기",
        "조교 평가는 우리 태스크 설정으로 로드하고 지형만 교체 → **학습 환경 + 로봇(센서·물성) + 학습 방법**이 설계 대상 (구조 변경은 불가)",
    ], size=16)  # fmt: skip
    notes(s, "연구 문제를 먼저 던지고, 진단 실험 두 개(성능 붕괴, 뒤집힘)로 문제를 좁혔다는 점을 강조.\n"
             "높이가 아니라 자세(뒤집힘)가 표적이라는 것이 이후 가설의 출발점.")  # fmt: skip

    # 3. evaluation design ------------------------------------------------------------------------------------------
    s = page(prs, 3, "실험 설계", "학습과 분리된 4단 평가", "출처: docs/report_material.md 4절")
    table(s, MX, Inches(1.8), CW, [
        ["단계", "지형", "용도"],
        ["분석용", "12종 — 평지 · 블록 6종 · 계단 · 파도 · 빙판 · 상자 · 레일", "가설 비교"],
        ["선택용 held-out", "Boxes · Rails", "체크포인트 선택에만"],
        ["보고용 held-out", "SlopedGrid · Pyramids", "선택 뒤 보고만"],
        ["잠금 테스트", "UnseenMix — 학습 범위 밖 블록 · 피라미드 · 경사 위 블록 · 원기둥 · 평지", "마지막 1회 (10.05)"],
    ], [2.3, 7.6, 2.4], size=15, hl=(4,), row_h=0.55)  # fmt: skip
    text(s, MX, Inches(4.85), CW, Inches(1.8), [
        "공식 명령 `play_one_episode.py --seed 24 --num_envs 100` · 원래 보상 · 첫 에피소드 누적 reward (조교 지표와 동일)",
        "학습 seed 2~4개 반복 · 한 번에 한 요인만 바꾸는 통제 비교",
        "선택 규칙 **\"선택용 held-out 1위\"를 사전 고정** → 고른 지형에 맞춘 결과를 막음",
    ], size=16)  # fmt: skip
    notes(s, "평가를 먼저 설계했다. 체크포인트를 고르는 지형과 결론을 말하는 지형을 분리해 \"고른 지형에 맞춘 결과\"를 막았다.\n"
             "잠금 테스트는 최종 선택이 끝난 뒤 한 번만 돌렸다.")  # fmt: skip

    # 4. approach ---------------------------------------------------------------------------------------------------
    s = page(prs, 4, "가설", "가설 5개, 세 갈래",
             "막대 = 학습 seed 평균, 점 = seed별 · 값 = 처음 보는 지형 12종 평균 reward")  # fmt: skip
    cols = [("축 1 · 환경", ["H1 경험 부족", "H2 자세 제약 부족"]),
            ("축 2 · 로봇", ["H3 감각 부족", "H4 접촉 물성"]),
            ("학습", ["H5 좌우 대칭 경험"])]  # fmt: skip
    cw3 = (CW - Inches(0.6)) / 3
    for i, (lab, hs) in enumerate(cols):
        cx = MX + i * (cw3 + Inches(0.3))
        hline(s, cx, Inches(1.75), cw3, "ink900", 1.5)
        kicker_line(s, cx, Inches(1.85), cw3, lab, "brand")
        text(s, cx, Inches(2.15), cw3, Inches(0.4), [dict(t="  ·  ".join(hs), size=16, face=SERIF, bold=True, gap=0)])
    fw = Inches(10.4)
    figure(s, "r2_ladder.png", MX + (CW - fw) / 2, Inches(2.65), fw, "변형별 처음 보는 지형 12종 평균 (docs/figures/r2_ladder.png)", 1)
    notes(s, "사다리 그림 하나로 전체를 보여 준다. 왼쪽(환경)은 40점대에서 멈추고, 로봇을 바꾸면 70점대, 대칭 증강으로 90점대.\n"
             "이후 슬라이드는 이 그림의 각 구간을 가설별로 설명.")  # fmt: skip

    # 5. axis 1 -----------------------------------------------------------------------------------------------------
    s = page(prs, 5, "H1 · H2 — 환경", "커리큘럼이 아니라 지형 비중, 자세 제약은 기각",
             "출처: results/summary_round2.md (최종 모델) · V6′·V10·V11은 seed 1개")  # fmt: skip
    tw = Inches(6.9)
    table(s, MX, Inches(1.8), tw, [
        ["실험", "12종 평균"],
        ["baseline (평면)", "16.3 (n=1)"],
        ["V1 지형 무작위", "44.4 ± 1.0 (n=2)"],
        ["V2 난이도 차선 커리큘럼", "47.5 ± 0.6 (n=3)"],
        ["A-noPromo (승급·강등 제거)", "46.8 ± 0.9 (n=2)"],
        ["A-randDiff (난이도 정렬 제거)", "47.9 ± 0.5 (n=2)"],
        ["V6′ 자세 제약 종료 CaT", "44.4 (n=1)"],
        ["V10 외란 + 기울기 페널티", "43.8 (n=1)"],
        ["V11 레벨별 지면 마찰", "37.5 (n=1)"],
    ], [4.4, 2.5], size=14, num_cols=(1,), hl=(3,), row_h=0.5)  # fmt: skip
    rx = MX + tw + Inches(0.5)
    rw = CW - tw - Inches(0.5)
    pill(s, rx, Inches(1.8), "H1 경험 부족 — 지지", "done", 15)
    text(s, rx, Inches(2.2), rw, Inches(1.6), [
        "지형을 경험시키면 16 → 47",
        "커리큘럼을 빼도 그대로 → **커리큘럼 효과 미미**",
        dict(t="핵심은 어떤 지형을 얼마나 경험하느냐", lv=2),
    ], size=16)  # fmt: skip
    pill(s, rx, Inches(3.9), "H2 자세 제약 부족 — 기각", "blocked", 15)
    text(s, rx, Inches(4.3), rw, Inches(2.0), [
        "자세를 제약하거나 흔들어도 모두 V2 이하",
        dict(t="평지 낙상률 V2 66% → V6′ 70% · V10 63%", lv=2),
        dict(t="조심시키면 덜 걸을 뿐, 뒤집힘은 줄지 않음", lv=2),
    ], size=16)  # fmt: skip
    notes(s, "H2는 팀원 의견에서 출발한 가설인데 실험으로 기각됐다. 조심시키면 덜 걸을 뿐 뒤집힘은 줄지 않았다 → 원인은 정보·능력 쪽이라고 보고 축 2로.\n"
             "V6′·V10·V11은 seed 1개라 근거가 약하다는 점도 말한다.")  # fmt: skip

    # 6. axis 2 -----------------------------------------------------------------------------------------------------
    s = page(prs, 6, "H3 · H4 — 로봇", "로봇을 바꾸는 쪽이 컸다", "출처: results/summary_round2.md (최종 모델)")
    tw = Inches(7.4)
    table(s, MX, Inches(1.8), tw, [
        ["실험", "12종 평균"],
        ["V2 (대조군)", "47.5 ± 0.6 (n=3)"],
        ["V12 발 접촉 센서 + 히스토리 3 step + 정규화", "53.1 ± 1.4 (n=2)"],
        ["V14 발 외피 재질 μ0.4", "56.0 ± 3.5 (n=2)"],
        ["V1214 = V12 + V14", "74.7 ± 1.5 (n=3)"],
        ["V1214B = V1214 + 블록 70% 지형", "78.3 ± 2.1 (n=4)"],
    ], [5.0, 2.4], size=14, num_cols=(1,), hl=(4,), best=((5, 1),), row_h=0.52)  # fmt: skip
    rx = MX + tw + Inches(0.5)
    rw = CW - tw - Inches(0.5)
    pill(s, rx, Inches(1.8), "H3 감각 부족 — 지지", "done", 15)
    pill(s, rx, Inches(2.2), "H4 접촉 물성 — 결합 시 지지", "accent", 15)
    callout(s, rx, Inches(2.75), rw, Inches(1.9), "감각 × 물성 상호작용", [
        dict(t="따로는 +5.6 / +8.5, 함께는 **+27.2**", size=16, gap=6),
        dict(t="감각이 외피의 미끄러짐을 보정한다는 해석 (직접 검증은 안 함)", size=13, color="ink700"),
    ])  # fmt: skip
    text(s, MX, Inches(5.15), CW, Inches(1.6), [
        "감각 분해 (V1214 기준, n=1): 히스토리 제거 68.5 · LSTM으로 대체 67.1 → **짧은 기억 + MLP**가 최선",
        "환경 쪽 변경(40점대)보다 **로봇 쪽 변경(70점대)** 이 컸다",
        dict(t="V12·V14·V1214는 V2 지형, V1214B만 블록 70% 지형 · 외피는 Python cfg만 수정 (물성치 변경 허용 공지)", lv=2),
    ], size=16)  # fmt: skip
    notes(s, "센서는 물리를 바꾸지 않고 정책이 아는 정보만 늘린다. 외피는 발과 지면의 마찰을 바꾼다.\n"
             "둘을 합쳤을 때 합보다 크게 오른 것이 핵심 — 감각이 외피의 미끄러짐을 보정한다는 해석 (직접 검증은 안 함).")  # fmt: skip

    # 7. mu sweep ---------------------------------------------------------------------------------------------------
    s = page(prs, 7, "H4 — 외피 마찰", "마찰에는 내부 최적이 있다", "출처: results/summary_round2.md, docs/log_1003.md")
    fw = Inches(6.3)
    figure(s, "r2_mu_sweep.png", MX, Inches(1.8), fw, "외피 마찰 sweep, V1214R 위 (docs/figures/r2_mu_sweep.png)", 2)
    rx = MX + fw + Inches(0.5)
    rw = CW - fw - Inches(0.5)
    table(s, rx, Inches(1.8), rw, [
        ["외피 μ", "12종 평균"],
        ["0.2", "71.5 (n=1)"],
        ["0.3", "73.6 ± 3.6 (n=3)"],
        ["0.4", "71.7 ± 7.2 (n=3)"],
        ["0.6", "67.7 (n=1)"],
    ], [1.2, 2.6], size=14, num_cols=(0, 1), best=((2, 1),), row_h=0.42)  # fmt: skip
    text(s, rx, Inches(4.0), rw, Inches(0.5), [dict(t="높으면 걸려 뒤집히고, 낮으면 미끄러진다 (예측과 일치)", size=14, color="ink700")])
    callout(s, rx, Inches(4.6), rw, Inches(1.95), "결론 철회 사례", [
        dict(t="seed 1개로 \"μ0.4가 μ0.3보다 낫다\" (79.4 vs 73.7, V1214B)", size=14, gap=4),
        dict(t="3 seed: 78.3 ± 2.5 = 78.3 ± 4.2 → **철회**", size=14, gap=4),
        dict(t="이후 결론은 seed 3개 이상에서만", size=13, color="ink700"),
    ], tone="accent")  # fmt: skip
    notes(s, "진단에서 본 \"마찰이 클수록 뒤집힘\"이 가설의 출발점. 외피 마찰을 바꿔 보니 중간값이 가장 좋았다.\n"
             "seed 하나로 결론을 냈다가 뒤집힌 경험을 그대로 보여 주는 것이 실험 설계의 신뢰성.")  # fmt: skip

    # 8. V8 ---------------------------------------------------------------------------------------------------------
    s = page(prs, 8, "H5 — 좌우 대칭", "거울 경험으로 +12.9",
             "코드: source/ant_rough/ant_rough/tasks/symmetry.py, scripts/check_symmetry.py · Mittal et al., ICRA 2024")  # fmt: skip
    lw = Inches(5.1)
    pill(s, MX, Inches(1.8), "H5 좌우 대칭 경험 — 지지", "done", 15)
    stats(s, MX, Inches(2.3), lw, [("91.2 ± 2.8", "V8 (n=4)", True), ("78.3 ± 2.1", "V1214B 대조군 (n=4)", False)], size=22)
    text(s, MX, Inches(3.55), lw, Inches(3.2), [
        "PPO 업데이트 때 상태·행동을 **좌우 거울로 뒤집어 추가** — rollout 1번으로 2배 경험, 파라미터 추가 없음",
        "4 seed 모두 V1214B 최고 seed보다 높음",
        "거울 맵을 **시뮬레이션에서 검증**: 거울 상태 쌍을 같이 굴려 비교 (240쌍, 상관 +0.81~1.00)",
        "발 반력 부호는 손 유도 실패 → **상관으로 측정**",
    ], size=15, gap=6)  # fmt: skip
    fx = MX + lw + Inches(0.45)
    figure(s, "r2_heatmap.png", fx, Inches(1.8), CW - lw - Inches(0.45),
           "지형별 reward — 맨 아래 V8 행이 거의 모든 지형에서 가장 높음 (r2_heatmap.png)", 3)  # fmt: skip
    notes(s, "이번 발표의 독창성 포인트. \"왼발로 익힌 요령은 오른발도 안다\"는 경험을 학습 데이터로 모델링.\n"
             "추측 대신 측정으로 맵을 확정한 과정을 짧게. 히트맵에서 맨 아래 V8 행이 거의 모든 지형에서 가장 진하다.")  # fmt: skip

    # 9. symmetry ablation ------------------------------------------------------------------------------------------
    s = page(prs, 9, "H5 — 어블레이션", "이득은 어디서 오는가", "출처: results/summary_round2.md, docs/log_1004.md")
    table(s, MX, Inches(1.8), CW, [
        ["실험", "12종 평균", "대조군 대비"],
        ["V1214B (대조군)", "78.3 ± 2.1 (n=4)", ""],
        ["V8 거울 데이터 증강", "91.2 ± 2.8 (n=4)", "+12.9"],
        ["V8-ML \"대칭으로 행동하라\" 손실만", "77.9 ± 1.1 (n=2)", "−0.4"],
        ["V2Sym 기본 로봇 + 증강 (대조군 V2 47.5)", "50.2 ± 2.7 (n=3)", "+2.7"],
    ], [6.2, 3.2, 2.6], size=16, num_cols=(1, 2), hl=(2,), best=((2, 2),), row_h=0.6)  # fmt: skip
    col = (CW - Inches(0.5)) / 2
    callout(s, MX, Inches(4.95), col, Inches(1.6), "대칭 강제 vs 거울 경험",
            [dict(t="손실만으로는 효과 없음 → 이득은 **거울 데이터(경험)** 에서", size=15)])  # fmt: skip
    callout(s, MX + col + Inches(0.5), Inches(4.95), col, Inches(1.6), "로봇 변경과의 상호작용", [
        dict(t="기본 로봇에서는 +2.7뿐 → **감각·물성 변경 위에서만 크게 작동**", size=15, gap=4),
        dict(t="해석 단계 (직접 검증은 안 함)", size=12, color="ink700"),
    ], tone="accent")  # fmt: skip
    notes(s, "대조 실험 두 개로 \"대칭인 정책\"과 \"거울 경험\"을 분리했다.\n"
             "감각이 풍부할 때 거울 데이터가 크게 작동한다는 상호작용은 해석 단계.")  # fmt: skip

    # 10. final -----------------------------------------------------------------------------------------------------
    s = page(prs, 10, "최종 검증", "잠금 테스트에서도 같은 순서",
             "출처: docs/report_material.md 6절, results/official_test.csv · 지형 밖 점검은 eval 스크립트 기준")  # fmt: skip
    stats(s, MX, Inches(1.75), CW, [
        ("v8_s44@2999", "최종 모델 · 선택용 held-out 1위 111.5", False),
        ("174.73 ± 12.67", "제출 태스크 공식 평면 (baseline 134.90)", False),
        ("113.65 ± 29.03", "잠금 테스트 UnseenMix (공식 조건, 1회)", True),
    ], size=22)  # fmt: skip
    table(s, MX, Inches(3.05), CW, [
        ["잠금 테스트 UnseenMix", "반영 가설", "reward", "평균 step (960 = 완주)"],
        ["baseline", "—", "9.51 ± 7.52", "761"],
        ["V2", "H1", "46.70 ± 30.17", "863"],
        ["V1214B", "H1 + H3 + H4", "100.19 ± 22.11", "946"],
        ["V8 (최종)", "+ H5", "113.65 ± 29.03", "938"],
    ], [3.8, 3.0, 2.9, 2.8], size=15, num_cols=(2, 3), hl=(4,), best=((4, 2),), row_h=0.45)  # fmt: skip
    text(s, MX, Inches(5.5), CW, Inches(1.3), [
        "자체 평가 순서가 처음 보는 잠금 테스트에서도 그대로 · 대칭 효과 +12.9 (자체) ↔ +13.5 (잠금)",
        "타당성 점검: 약 27%가 평가 지형 밖으로 나감 → 지형 안 로봇만 비교해도 순서 그대로 (9.1 < 42.6 < 87.8 < 104.6)",
    ], size=15, gap=5)  # fmt: skip
    notes(s, "자체 평가에서 본 순서가 처음 보는 잠금 테스트에서도 그대로 나왔다 → 결론이 평가 지형에 과적합되지 않았다.\n"
             "평가 지형 이탈은 스스로 찾아 점검했다는 점을 말한다.")  # fmt: skip

    # 11. videos ----------------------------------------------------------------------------------------------------
    s = page(prs, 11, "영상", "최종 모델이 걷는 모습",
             "docs/media/ (git 미포함, 워크스테이션) · 1280×720 · 960 step = 완주 · 목록 docs/media/videos.csv")  # fmt: skip
    clips = [("평면 (제출 태스크)", "final_plane.mp4", "180.5", "960"), ("평지 (메시)", "final_flat.mp4", "144.0", "960"),
             ("레일", "final_rails.mp4", "137.9", "960"), ("블록", "final_grid.mp4", "125.4", "960"),
             ("피라미드", "final_pyramids.mp4", "105.8", "960"), ("계단", "final_stairs.mp4", "101.9", "960"),
             ("평지 seed 27 — 뒤집힘", "final_flat_fail.mp4", "100.7", "713")]  # fmt: skip
    gap = Inches(0.25)
    bw = (CW - 3 * gap) / 4
    bh = bw * 9 / 16
    for i, (t, f, r, st) in enumerate(clips):
        row, c = divmod(i, 4)
        x = MX + c * (bw + gap)
        y = Inches(1.7) + row * (bh + Inches(1.05))
        fail = "fail" in f
        rect(s, x, y, bw, bh + Inches(0.88), fill="white", line="blocked" if fail else "ink200")
        box = rect(s, x, y, bw, bh, fill="ink900")
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        runs(box.text_frame.paragraphs[0], "영상 자리", 11, color="ink500")
        text(s, x + Inches(0.12), y + bh + Inches(0.08), bw - Inches(0.24), Inches(0.7), [
            dict(t=t, size=13, color="blocked" if fail else "ink900", bold=True, gap=2),
            dict(t=f, size=10, face=MONO, color="data", gap=0),
            dict(t=f"reward {r} · {st} step", size=10, face=MONO, color="ink600", gap=0),
        ])  # fmt: skip
    notes(s, "블록·계단 영상 하나와 실패 영상 하나를 보여 준다. 실패도 713 step까지 걸은 뒤 뒤집힘 — 남은 한계로 연결.\n"
             "\"영상 자리\" 상자에 영상(또는 첫 프레임)을 넣기.")  # fmt: skip

    # 12. limits ----------------------------------------------------------------------------------------------------
    s = page(prs, 12, "한계", "한계와 배운 것")
    text(s, MX, Inches(1.8), CW, Inches(3.9), [
        "**뒤집힘 자체는 남았다**: 평지 낙상률 V2 66.3% / V8 65.8%",
        dict(t="향상은 넘어지기 전까지 더 멀리 걷는 데서 왔다 → 다음 연구 문제", lv=2),
        "평가 지형 폭(32 m)이 빠른 정책에는 좁다: 약 27% 이탈 → 결론은 같지만 넓은 지형 재평가는 못 함",
        "높이 스캔(RayCaster) 미사용: 정적 메시 1개만 지원 → 조교 지형이 박스 prim이면 에러",
        dict(t="지형과 무관한 발 접촉 센서 + 히스토리로 대체", lv=2),
        "seed 1개 결론은 뒤집힐 수 있다 (μ 철회) → 결론은 seed 3개 이상 · V6′·V10·V11은 seed 1개",
        "실험 운영: 백그라운드 실행 알림 누락으로 GPU가 놀았던 일 → 실행 도구·절대 경로로 통일",
    ], size=17, gap=8)  # fmt: skip
    callout(s, MX, Inches(5.6), CW, Inches(1.15), "한 줄 요약",
            [dict(t="처음 보는 지형에서 걷게 만든 건 \"조심\"이 아니라 감각과 경험이었다", size=19, face=SERIF, bold=True, color="brand")])  # fmt: skip
    notes(s, "한계를 숨기지 않는다. 특히 \"뒤집힘은 그대로\"는 다음 연구 문제로 제시.\n"
             "참고문헌: Rudin CoRL 2021 · Cobbe ICML 2019 · Chane-Sane (CaT) IROS 2024 · Kumar (RMA) RSS 2021 · "
             "Mittal ICRA 2024")  # fmt: skip

    prs.save(OUT)
    print(f"[SLIDES] {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
