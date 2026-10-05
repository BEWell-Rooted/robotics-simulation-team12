"""Draft presentation deck (v0) for assignment 1 -> docs/assignment1_slides_v0.pptx.

Text mirrors docs/slides_draft_v0.md (numbers from results/summary_round2.md and docs/report_material.md).
16:9, 12 slides, figures from docs/figures/, videos only as grey placeholders with their paths.

    <venv with python-pptx>/bin/python scripts/make_slides_v0.py
"""

import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PROJECT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(PROJECT, "docs", "figures")
OUT = os.path.join(PROJECT, "docs", "assignment1_slides_v0.pptx")

FONT = "맑은 고딕"  # Noto Sans KR is not installed on the workstation (see slides_draft_v0.md)
INK, INK2, MUTED = RGBColor(0x1A, 0x1A, 0x1A), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x8A, 0x88, 0x84)
ACCENT, ACCENT_BG = RGBColor(0x6B, 0x4B, 0xC4), RGBColor(0xEE, 0xEA, 0xF8)
PANEL, LINE = RGBColor(0xF3, 0xF2, 0xEF), RGBColor(0xD9, 0xD7, 0xD2)
OK, NO = RGBColor(0x1B, 0x8A, 0x5A), RGBColor(0xC0, 0x39, 0x2B)
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)


def set_font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def add_runs(par, text, size, color=INK, bold=False):
    """Text with **bold** spans."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            set_font(par.add_run(), size, bold=bold or i % 2 == 1, color=color)
            par.runs[-1].text = part


def textbox(slide, left, top, width, height, lines, size=16, color=INK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT,
            spacing=6):  # fmt: skip
    """lines: str or (text, level) / (text, level, size, color, bold). level 0 = plain, 1 = bullet, 2 = sub-bullet."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, line in enumerate(lines):
        if isinstance(line, str):
            line = (line, 1)
        text, level = line[0], line[1]
        sz = line[2] if len(line) > 2 else size - (2 if level == 2 else 0)
        col = line[3] if len(line) > 3 else (INK2 if level == 2 else color)
        bold = line[4] if len(line) > 4 else False
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.space_after = Pt(spacing)
        bullet = {0: "", 1: "• ", 2: "– "}[level]
        if level == 2:
            par.level = 1
        add_runs(par, bullet + text, sz, color=col, bold=bold)
    return tb


def title(slide, text, sub=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    textbox(slide, MARGIN, Inches(0.35), W - 2 * MARGIN, Inches(0.8), [(text, 0, 28, INK, True)])
    if sub:
        textbox(slide, MARGIN, Inches(1.05), W - 2 * MARGIN, Inches(0.5), [(sub, 0, 15, INK2)])


def footer(slide, n, source=None):
    textbox(slide, W - Inches(1.2), H - Inches(0.5), Inches(0.8), Inches(0.35), [(f"{n} / 12", 0, 10, MUTED)],
            align=PP_ALIGN.RIGHT)  # fmt: skip
    if source:
        textbox(slide, MARGIN, H - Inches(0.5), Inches(10), Inches(0.35), [(source, 0, 10, MUTED)])


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def table(slide, left, top, width, rows, col_w, size=13, highlight=(), row_h=0.38):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, Inches(row_h * len(rows)))
    tbl = shape.table
    total = sum(col_w)
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Emu(int(width * w / total))
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else (ACCENT_BG if i in highlight else RGBColor(0xFF, 0xFF, 0xFF))
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            add_runs(par, val, size, color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK, bold=i == 0 or i in highlight)
    return shape


def picture(slide, name, left, top, width=None, height=None):
    return slide.shapes.add_picture(os.path.join(FIG, name), left, top, width=width, height=height)


def panel(slide, left, top, width, height, fill=PANEL, line=None):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.adjustments[0] = 0.08
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
    return box


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    new = lambda: prs.slides.add_slide(blank)  # noqa: E731

    # 1. cover ------------------------------------------------------------------------------------------------------
    s = new()
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.5), H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = ACCENT
    bg.line.fill.background()
    textbox(s, Inches(1.2), Inches(2.2), Inches(11), Inches(1.2), [("처음 보는 지형에서도 걷는 Ant", 0, 44, INK, True)])
    textbox(s, Inches(1.2), Inches(3.4), Inches(11), Inches(0.6), [("로보틱스 시뮬레이션 과제 1 · 팀 12", 0, 22, INK2)])
    textbox(s, Inches(1.2), Inches(4.6), Inches(11), Inches(1.0),
            [("박형근 · 경한 · [구성원]", 0, 18, INK), ("2026.10.08", 0, 16, MUTED)])  # fmt: skip
    notes(s, "팀 소개. \"평면에서만 배운 Ant를 처음 보는 지형에서도 걷게 만든 과정을 가설과 검증 순서로 말씀드리겠습니다.\"\n"
             "[구성원] 이름 채우기.")  # fmt: skip

    # 2. problem ----------------------------------------------------------------------------------------------------
    s = new()
    title(s, "문제 — 왜 처음 보는 지형에서 무너지는가")
    panel(s, MARGIN, Inches(1.35), W - 2 * MARGIN, Inches(0.95), fill=ACCENT_BG)
    textbox(s, MARGIN + Inches(0.3), Inches(1.4), W - 2 * MARGIN - Inches(0.6), Inches(0.85),
            [("연구 문제: 평면에서 학습한 Ant는 왜 처음 보는 지형에서 무너지는가?", 0, 22, ACCENT, True)],
            anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    textbox(s, MARGIN, Inches(2.6), W - 2 * MARGIN, Inches(4.2), [
        "과제: Isaac-Ant-v0 정책이 **지형 형태·마찰만 바뀐 처음 보는 환경**에서도 걷게 만들기",
        "관찰 1 — **성능 붕괴**: baseline 공식 평면 134.90 ± 27.74 → 처음 보는 지형 12종 평균 **16.3**",
        "관찰 2 — **실패 = 뒤집힘**: 종료 높이 임계를 0.31 → 0.20으로 낮춰도 전진 거리 그대로 (59 m)",
        ("이미 뒤집혀 누운 상태 → 표적은 높이가 아니라 자세", 2),
        "조교 평가: 우리 태스크 설정으로 로드하고 지형만 교체",
        ("**학습 환경 + 로봇(센서·물성) + 학습 방법**이 설계 대상 (구조 변경은 불가, 10.05 공지)", 2),
    ], size=18)  # fmt: skip
    footer(s, 2, "출처: docs/report_material.md 2절 (진단 실험 log_0930)")
    notes(s, "연구 문제를 먼저 던지고, 진단 실험 두 개(성능 붕괴, 뒤집힘)로 문제를 좁혔다는 점을 강조.\n"
             "높이가 아니라 자세(뒤집힘)가 표적이라는 것이 이후 가설의 출발점.")  # fmt: skip

    # 3. evaluation design ------------------------------------------------------------------------------------------
    s = new()
    title(s, "평가 설계 — 학습과 분리된 4단 평가")
    table(s, MARGIN, Inches(1.4), W - 2 * MARGIN, [
        ["단계", "지형", "용도"],
        ["분석용", "12종 (평지·블록 6종·계단·파도·빙판·상자·레일)", "가설 비교"],
        ["선택용 held-out", "Boxes · Rails", "체크포인트 선택에만"],
        ["보고용 held-out", "SlopedGrid · Pyramids", "선택 뒤 보고만"],
        ["잠금 테스트", "UnseenMix (학습 범위 밖 블록·피라미드·경사 위 블록·원기둥·평지)", "마지막 1회 (10.05)"],
    ], [2.2, 7.5, 2.8], size=15, highlight=(4,), row_h=0.55)  # fmt: skip
    textbox(s, MARGIN, Inches(4.5), W - 2 * MARGIN, Inches(2.3), [
        "공식 명령 play_one_episode.py --seed 24 --num_envs 100, 원래 보상, 첫 에피소드 누적 reward (조교 지표와 동일)",
        "학습 seed 2~4개 반복 · 한 번에 한 요인만 바꾸는 비교",
        "선택 규칙 **\"선택용 held-out 1위\"를 사전 고정** → 고른 지형에 맞춘 결과를 막음",
    ], size=17)  # fmt: skip
    footer(s, 3, "출처: docs/report_material.md 4절")
    notes(s, "평가를 먼저 설계했다. 체크포인트를 고르는 지형과 결론을 말하는 지형을 분리해 \"고른 지형에 맞춘 결과\"를 막았다.\n"
             "잠금 테스트는 최종 선택이 끝난 뒤 한 번만 돌렸다.")  # fmt: skip

    # 4. approach ---------------------------------------------------------------------------------------------------
    s = new()
    title(s, "접근 한눈에 — 가설 5개, 세 갈래")
    chips = [("축 1 환경", "H1 경험 부족 · H2 자세 제약 부족", RGBColor(0x2A, 0x78, 0xD6)),
             ("축 2 로봇", "H3 감각 부족 · H4 접촉 물성", RGBColor(0xEB, 0x68, 0x34)),
             ("학습", "H5 좌우 대칭 경험", ACCENT)]  # fmt: skip
    cw = (W - 2 * MARGIN - Inches(0.4)) / 3
    for i, (head, body, col) in enumerate(chips):
        x = MARGIN + i * (cw + Inches(0.2))
        panel(s, x, Inches(1.3), cw, Inches(0.95), fill=PANEL, line=col)
        textbox(s, x + Inches(0.15), Inches(1.33), cw - Inches(0.3), Inches(0.9),
                [(head, 0, 15, col, True), (body, 0, 15, INK)], spacing=2, anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    pic_w = W - 2 * MARGIN
    picture(s, "r2_ladder.png", MARGIN, Inches(2.45), width=pic_w)
    footer(s, 4, "그림: docs/figures/r2_ladder.png — 막대 = 학습 seed 평균, 점 = seed별, 자체 평가 12종 평균 reward")
    notes(s, "사다리 그림 하나로 전체를 보여 준다. 왼쪽(환경)은 40점대에서 멈추고, 로봇을 바꾸면 70점대, 대칭 증강으로 90점대.\n"
             "이후 슬라이드는 이 그림의 각 구간을 가설별로 설명.")  # fmt: skip

    # 5. axis 1 -----------------------------------------------------------------------------------------------------
    s = new()
    title(s, "축 1 결과 — H1 지지 (커리큘럼은 기각, 지형 비중이 핵심) · H2 기각")
    table(s, MARGIN, Inches(1.4), Inches(6.6), [
        ["실험", "12종 평균"],
        ["baseline (평면)", "16.3 (n=1)"],
        ["V1 지형 무작위", "44.4 ± 1.0 (n=2)"],
        ["V2 난이도 차선 커리큘럼", "47.5 ± 0.6 (n=3)"],
        ["A-noPromo (승급·강등 제거)", "46.8 ± 0.9 (n=2)"],
        ["A-randDiff (난이도 정렬 제거)", "47.9 ± 0.5 (n=2)"],
        ["V6′ 자세 제약 종료 CaT", "44.4 (n=1)"],
        ["V10 외란 + 기울기 페널티", "43.8 (n=1)"],
        ["V11 레벨별 지면 마찰", "37.5 (n=1)"],
    ], [4.2, 2.4], size=14, highlight=(3,), row_h=0.5)  # fmt: skip
    textbox(s, Inches(7.6), Inches(1.4), Inches(5.2), Inches(5.2), [
        ("H1 경험 부족 — 지지", 0, 18, OK, True),
        "지형을 경험시키면 16 → 47",
        "커리큘럼을 빼도 그대로 → **커리큘럼 효과 미미**",
        ("핵심은 어떤 지형을 얼마나 경험하느냐", 2),
        ("H2 자세 제약 부족 — 기각", 0, 18, NO, True),
        "자세를 제약하거나 흔들어도 모두 V2 이하",
        ("평지 낙상률 V2 66% → V6′ 70% · V10 63%", 2),
        ("조심시키면 덜 걸을 뿐, 뒤집힘은 줄지 않음", 2),
    ], size=16)  # fmt: skip
    footer(s, 5, "출처: results/summary_round2.md (최종 모델) · V6′·V10·V11은 seed 1개")
    notes(s, "H2는 팀원 의견에서 출발한 가설인데 실험으로 기각됐다. 조심시키면 덜 걸을 뿐 뒤집힘은 줄지 않았다 → 원인은 정보·능력 쪽이라고 보고 축 2로.\n"
             "V6′·V10·V11은 seed 1개라 근거가 약하다는 점도 말한다.")  # fmt: skip

    # 6. axis 2 -----------------------------------------------------------------------------------------------------
    s = new()
    title(s, "축 2 결과 — H3 지지 · H4 결합 시 지지: 로봇을 바꾸는 쪽이 컸다")
    table(s, MARGIN, Inches(1.4), Inches(7.4), [
        ["실험", "12종 평균"],
        ["V2 (대조군)", "47.5 ± 0.6 (n=3)"],
        ["V12 발 접촉 센서 + 히스토리 3 step + 정규화", "53.1 ± 1.4 (n=2)"],
        ["V14 발 외피 재질 μ0.4", "56.0 ± 3.5 (n=2)"],
        ["V1214 = V12 + V14", "74.7 ± 1.5 (n=3)"],
        ["V1214B = V1214 + 블록 70% 지형", "78.3 ± 2.1 (n=4)"],
    ], [5.0, 2.4], size=14, highlight=(4, 5), row_h=0.55)  # fmt: skip
    textbox(s, Inches(8.4), Inches(1.4), Inches(4.4), Inches(3.4), [
        ("감각 × 물성 상호작용", 0, 18, ACCENT, True),
        "따로는 +5.6 / +8.5",
        "함께는 **+27.2** (합 14.1보다 큼)",
        ("감각이 외피의 미끄러짐을 보정한다는 해석 (직접 검증은 안 함)", 2),
    ], size=16)  # fmt: skip
    textbox(s, MARGIN, Inches(5.0), W - 2 * MARGIN, Inches(1.8), [
        "감각 분해 (V1214 기준, n=1): 히스토리 제거 68.5 · LSTM으로 대체 67.1 → **짧은 기억 + MLP**가 최선",
        "환경 쪽 변경(40점대)보다 **로봇 쪽 변경(70점대)** 이 컸다",
        ("V12·V14·V1214는 V2 지형, V1214B만 블록 70% 지형 · 외피는 Python cfg만 수정 (물성치 변경 허용 공지)", 2),
    ], size=16)  # fmt: skip
    footer(s, 6, "출처: results/summary_round2.md (최종 모델)")
    notes(s, "센서는 물리를 바꾸지 않고 정책이 아는 정보만 늘린다. 외피는 발과 지면의 마찰을 바꾼다.\n"
             "둘을 합쳤을 때 합보다 크게 오른 것이 핵심 — 감각이 외피의 미끄러짐을 보정한다는 해석 (직접 검증은 안 함).")  # fmt: skip

    # 7. mu sweep ---------------------------------------------------------------------------------------------------
    s = new()
    title(s, "외피 마찰 sweep — H4: 마찰에는 내부 최적이 있다")
    picture(s, "r2_mu_sweep.png", MARGIN, Inches(1.35), width=Inches(6.4))
    textbox(s, Inches(7.3), Inches(1.4), Inches(5.5), Inches(2.6), [
        ("V1214R 위 μ sweep (12종 평균)", 0, 17, INK, True),
        "μ 0.2 — 71.5 (n=1)",
        "μ 0.3 — **73.6 ± 3.6 (n=3)**",
        "μ 0.4 — 71.7 ± 7.2 (n=3)",
        "μ 0.6 — 67.7 (n=1)",
        ("높으면 걸려 뒤집히고, 낮으면 미끄러진다 (예측과 일치)", 2),
    ], size=16, spacing=4)  # fmt: skip
    panel(s, Inches(7.3), Inches(4.35), Inches(5.45), Inches(2.3), fill=ACCENT_BG)
    textbox(s, Inches(7.45), Inches(4.4), Inches(5.2), Inches(2.2), [
        ("결론 철회 사례", 0, 17, ACCENT, True),
        "V1214B에서 seed 1개로 \"μ0.4가 μ0.3보다 낫다\" (79.4 vs 73.7)",
        "3 seed: μ0.4 78.3 ± 2.5 = μ0.3 78.3 ± 4.2 → **철회**",
        ("이후 결론은 seed 3개 이상에서만", 2),
    ], size=15, spacing=4)  # fmt: skip
    footer(s, 7, "그림: docs/figures/r2_mu_sweep.png · 출처: summary_round2.md, log_1003.md")
    notes(s, "진단에서 본 \"마찰이 클수록 뒤집힘\"이 가설의 출발점. 외피 마찰을 바꿔 보니 중간값이 가장 좋았다.\n"
             "seed 하나로 결론을 냈다가 뒤집힌 경험을 그대로 보여 주는 것이 실험 설계의 신뢰성.")  # fmt: skip

    # 8. V8 ---------------------------------------------------------------------------------------------------------
    s = new()
    title(s, "대칭 증강 V8 — H5 지지: 거울 경험으로 +12.9")
    textbox(s, MARGIN, Inches(1.35), Inches(5.3), Inches(5.4), [
        "Ant 몸은 좌우 대칭 → PPO 업데이트 때 상태·행동을 **좌우 거울로 뒤집어 추가**",
        ("rollout 1번으로 2배 경험, 파라미터 추가 없음 (Mittal et al., ICRA 2024)", 2),
        "**V8 91.2 ± 2.8 (n=4)** vs V1214B 78.3 ± 2.1 (n=4)",
        ("4 seed 모두 V1214B 최고 seed보다 높음", 2),
        "거울 맵을 **시뮬레이션에서 검증**",
        ("env 1을 env 0의 거울 상태로 두고 거울 행동으로 같이 굴려 비교 (240쌍, 상관 +0.81~1.00)", 2),
        "발 반력(관절 wrench) 부호는 손 유도 실패 → **상관으로 측정** (0.81~0.99)",
    ], size=16)  # fmt: skip
    picture(s, "r2_heatmap.png", Inches(6.15), Inches(1.4), width=Inches(6.6))
    footer(s, 8, "그림: docs/figures/r2_heatmap.png (지형별 reward) · 코드: tasks/symmetry.py, scripts/check_symmetry.py")
    notes(s, "이번 발표의 독창성 포인트. \"왼발로 익힌 요령은 오른발도 안다\"는 경험을 학습 데이터로 모델링.\n"
             "추측 대신 측정으로 맵을 확정한 과정을 짧게. 히트맵에서 맨 아래 V8 행이 거의 모든 지형에서 가장 진하다.")  # fmt: skip

    # 9. symmetry ablation ------------------------------------------------------------------------------------------
    s = new()
    title(s, "대칭 어블레이션 — 이득은 어디서 오는가")
    table(s, MARGIN, Inches(1.4), W - 2 * MARGIN, [
        ["실험", "12종 평균", "대조군 대비"],
        ["V1214B (대조군)", "78.3 ± 2.1 (n=4)", ""],
        ["V8 거울 데이터 증강", "91.2 ± 2.8 (n=4)", "+12.9"],
        ["V8-ML \"대칭으로 행동하라\" 손실만", "77.9 ± 1.1 (n=2)", "−0.4"],
        ["V2Sym 기본 로봇 + 증강 (대조군 V2 47.5)", "50.2 ± 2.7 (n=3)", "+2.7"],
    ], [6.0, 3.2, 2.8], size=16, highlight=(2,), row_h=0.6)  # fmt: skip
    textbox(s, MARGIN, Inches(4.75), W - 2 * MARGIN, Inches(2.0), [
        "대칭을 강제하는 손실만으로는 효과 없음 → 이득은 **거울 데이터(경험)** 에서",
        "감각·물성 변경이 없는 기본 로봇에서는 +2.7뿐 → **로봇 변경 위에서만 크게 작동** (+12.9 vs +2.7)",
        ("감각이 풍부할 때 거울 데이터가 크게 작동한다는 상호작용은 해석 단계 (직접 검증은 안 함)", 2),
    ], size=17)  # fmt: skip
    footer(s, 9, "출처: results/summary_round2.md, docs/log_1004.md")
    notes(s, "대조 실험 두 개로 \"대칭인 정책\"과 \"거울 경험\"을 분리했다.\n"
             "감각이 풍부할 때 거울 데이터가 크게 작동한다는 상호작용은 해석 단계.")  # fmt: skip

    # 10. final -----------------------------------------------------------------------------------------------------
    s = new()
    title(s, "최종 제출과 잠금 테스트")
    textbox(s, MARGIN, Inches(1.3), W - 2 * MARGIN, Inches(1.1), [
        "최종 모델 **v8_s44@2999**: 사전 규칙(선택용 held-out 1위) 111.5 · 낙상률 0.117 · 보고용 98.4",
        "제출 태스크 Isaac-Ant-Team12-v0, 공식 평면 **174.73 ± 12.67** (baseline 134.90 ± 27.74)",
    ], size=16, spacing=3)  # fmt: skip
    table(s, MARGIN, Inches(2.5), W - 2 * MARGIN, [
        ["잠금 테스트 UnseenMix (공식 조건, 1회)", "반영 가설", "reward", "평균 step (960 = 완주)"],
        ["baseline", "—", "9.51 ± 7.52", "761"],
        ["V2", "H1", "46.70 ± 30.17", "863"],
        ["V1214B", "H1 + H3 + H4", "100.19 ± 22.11", "946"],
        ["V8 (최종)", "+ H5", "113.65 ± 29.03", "938"],
    ], [4.2, 2.6, 2.8, 2.9], size=15, highlight=(4,), row_h=0.5)  # fmt: skip
    textbox(s, MARGIN, Inches(5.25), W - 2 * MARGIN, Inches(1.6), [
        "자체 평가 순서가 처음 보는 잠금 테스트에서도 그대로 · 대칭 효과 +12.9 (자체) ↔ +13.5 (잠금)",
        "타당성 점검: 약 27%가 평가 지형 밖으로 나감 → 지형 안 로봇만 비교해도 순서 그대로 (9.1 < 42.6 < 87.8 < 104.6)",
    ], size=15, spacing=4)  # fmt: skip
    footer(s, 10, "출처: docs/report_material.md 6절, results/official_test.csv · 지형 밖 점검은 eval 스크립트 기준")
    notes(s, "자체 평가에서 본 순서가 처음 보는 잠금 테스트에서도 그대로 나왔다 → 결론이 평가 지형에 과적합되지 않았다.\n"
             "평가 지형 이탈은 스스로 찾아 점검했다는 점을 말한다.")  # fmt: skip

    # 11. videos ----------------------------------------------------------------------------------------------------
    s = new()
    title(s, "영상 — 최종 모델", "docs/media/ (git 미포함, 워크스테이션) · 1280×720 · 960 step = 완주")
    clips = [("final_plane.mp4", "평면 (제출 태스크)", "180.5 · 960"), ("final_flat.mp4", "평지 (메시)", "144.0 · 960"),
             ("final_rails.mp4", "레일", "137.9 · 960"), ("final_grid.mp4", "블록", "125.4 · 960"),
             ("final_pyramids.mp4", "피라미드", "105.8 · 960"), ("final_stairs.mp4", "계단", "101.9 · 960"),
             ("final_flat_fail.mp4", "평지 seed 27 — 뒤집힘", "100.7 · 713")]  # fmt: skip
    bw, bh, gap = Inches(2.85), Inches(1.6), Inches(0.2)
    for i, (f, t, r) in enumerate(clips):
        row, col = divmod(i, 4)
        x = MARGIN + col * (bw + gap) + (0 if row == 0 else (bw + gap) / 2)
        y = Inches(1.75) + row * (bh + Inches(0.95))
        box = panel(s, x, y, bw, bh, fill=RGBColor(0xC8, 0xC6, 0xC2) if "fail" not in f else RGBColor(0xE8, 0xC4, 0xBE))
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_runs(box.text_frame.paragraphs[0], "썸네일", 12, color=INK2)
        textbox(s, x, y + bh + Inches(0.05), bw, Inches(0.8),
                [(t, 0, 13, INK, True), (f"{f} · {r.split(' · ')[0]} · {r.split(' · ')[1]} step", 0, 10, INK2)],
                spacing=1)  # fmt: skip
    footer(s, 11, "출처: docs/media/videos.csv")
    notes(s, "블록·계단 영상 하나와 실패 영상 하나를 보여 준다. 실패도 713 step까지 걸은 뒤 뒤집힘 — 남은 한계로 연결.\n"
             "썸네일 자리에 영상 첫 프레임 또는 영상을 넣기.")  # fmt: skip

    # 12. limits ----------------------------------------------------------------------------------------------------
    s = new()
    title(s, "한계와 배운 것")
    textbox(s, MARGIN, Inches(1.35), W - 2 * MARGIN, Inches(4.6), [
        "**뒤집힘 자체는 남았다**: 평지 낙상률 V2 66.3% / V8 65.8%",
        ("향상은 넘어지기 전까지 더 멀리 걷는 데서 왔다 → 다음 연구 문제", 2),
        "평가 지형 폭(32 m)이 빠른 정책에는 좁다: 약 27% 이탈 → 결론은 같지만 넓은 지형 재평가는 못 함",
        "높이 스캔(RayCaster) 미사용: 정적 메시 1개만 지원 → 조교 지형이 박스 prim이면 에러",
        ("지형과 무관한 발 접촉 센서 + 히스토리로 대체", 2),
        "seed 1개 결론은 뒤집힐 수 있다 (μ 철회) → 결론은 seed 3개 이상 · V6′·V10·V11은 seed 1개",
        "실험 운영: 백그라운드 실행 알림 누락으로 GPU가 놀았던 일 → 실행 도구·절대 경로로 통일",
    ], size=17)  # fmt: skip
    panel(s, MARGIN, Inches(6.05), W - 2 * MARGIN, Inches(0.75), fill=ACCENT_BG)
    textbox(s, MARGIN + Inches(0.3), Inches(6.08), W - 2 * MARGIN - Inches(0.6), Inches(0.7),
            [("처음 보는 지형에서 걷게 만든 건 \"조심\"이 아니라 감각과 경험이었다", 0, 18, ACCENT, True)],
            anchor=MSO_ANCHOR.MIDDLE)  # fmt: skip
    footer(s, 12)
    notes(s, "한계를 숨기지 않는다. 특히 \"뒤집힘은 그대로\"는 다음 연구 문제로 제시.\n"
             "참고문헌: Rudin CoRL 2021 · Cobbe ICML 2019 · Chane-Sane (CaT) IROS 2024 · Kumar (RMA) RSS 2021 · "
             "Mittal ICRA 2024")  # fmt: skip

    prs.save(OUT)
    print(f"[SLIDES] {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
